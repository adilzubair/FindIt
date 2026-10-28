# /// script
# requires-python = ">=3.11,<3.13"
# dependencies = [
#   "pillow==11.3.0",
#   "python-docx==1.2.0",
#   "openpyxl==3.1.5",
#   "pypdfium2==4.30.0",
#   "kokoro-onnx==0.4.9",
#   "soundfile==0.13.1",
#   "numpy<2.3",
# ]
# ///
"""Builds the FindIt demo drive: photos, videos, voice memos and documents.

    uv run demo-video/library/build.py --out "%USERPROFILE%\\FindItDemo\\Portable SSD"

Downloads are pinned by SHA-256 in lock.json (refresh with --update-lock). The output
folder is created from scratch and never contains anything but generated files.
"""
from __future__ import annotations

import argparse, datetime as dt, hashlib, json, os, random, shutil, subprocess, sys, tempfile, urllib.error, urllib.request
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parent.parent
UA = {"User-Agent": "Mozilla/5.0 (FindIt demo builder)"}
KOKORO = "https://github.com/thewh1teagle/kokoro-onnx/releases/download/model-files-v1.0/"
# Each part and the top-level folders it owns, so a partial rebuild replaces only its own files.
PARTS = {"photos": ["Photos"], "videos": ["Videos"], "memos": ["Voice memos"], "documents": ["Documents"]}


def log(*a):
    print(*a, flush=True)


class Downloads:
    def __init__(self, cache: Path, update_lock: bool):
        self.cache, self.update = cache, update_lock
        self.lock_path = HERE / "lock.json"
        self.lock = json.loads(self.lock_path.read_text()) if self.lock_path.exists() else {}

    def get(self, key: str, url: str) -> Path:
        path = self.cache / key
        if not path.exists():
            path.parent.mkdir(parents=True, exist_ok=True)
            log(f"  download {url}")
            tmp = path.with_suffix(path.suffix + ".partial")
            with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=120) as r, open(tmp, "wb") as f:
                shutil.copyfileobj(r, f, 1 << 20)
            tmp.replace(path)
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        expected = self.lock.get(key)
        if expected is None:
            if not self.update:
                sys.exit(f"{key} is not pinned in lock.json. Run with --update-lock to record it.")
            self.lock[key] = digest
        elif expected != digest:
            path.replace(path.with_suffix(path.suffix + ".rejected"))  # replace, not rename: an earlier .rejected may be there
            sys.exit(f"SHA-256 mismatch for {key}: expected {expected}, got {digest}")
        return path

    def save(self):
        if self.update:
            self.lock_path.write_text(json.dumps(dict(sorted(self.lock.items())), indent=2) + "\n")

    def clip(self, clip_id: int) -> Path:
        # Mixkit publishes different renditions per clip; take the first that exists, preferring 1080p.
        for res in (1080, 2160, 720):
            key = f"mixkit/{clip_id}-{res}.mp4"
            if (self.cache / key).exists() or key in self.lock:
                return self.get(key, f"https://assets.mixkit.co/videos/{clip_id}/{clip_id}-{res}.mp4")
        for res in (1080, 2160, 720):
            try:
                return self.get(f"mixkit/{clip_id}-{res}.mp4", f"https://assets.mixkit.co/videos/{clip_id}/{clip_id}-{res}.mp4")
            except urllib.error.HTTPError as e:
                if e.code not in (403, 404):
                    raise
        sys.exit(f"Mixkit clip {clip_id} has no downloadable rendition.")


def tools():
    ff = os.environ.get("FFMPEG") or str(REPO / "runtime" / "tools" / "ffmpeg.exe")
    if not Path(ff).exists():
        sys.exit("ffmpeg not found. Set FFMPEG or run scripts/Fetch-Dependencies.ps1.")
    probe = str(Path(ff).with_name("ffprobe.exe"))
    # Chrome first: managed Edge installs can silently refuse --print-to-pdf.
    edge = os.environ.get("BROWSER") or next((p for p in (
        r"C:\Program Files\Google\Chrome\Application\chrome.exe",
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe") if Path(p).exists()), None)
    return ff, probe, edge


def run(*args):
    subprocess.run([str(a) for a in args], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)


def duration(probe, path) -> float:
    out = subprocess.run([probe, "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", str(path)], capture_output=True, text=True, check=True)
    return float(out.stdout.strip())


def stamp(path: Path, when: str):
    t = dt.datetime.fromisoformat(when).timestamp()
    os.utime(path, (t, t))


def video_encoder(ff) -> list[str]:
    encoders = subprocess.run([ff, "-hide_banner", "-encoders"], capture_output=True, text=True).stdout
    if "h264_nvenc" in encoders:
        probe = subprocess.run([ff, "-v", "error", "-f", "lavfi", "-i", "color=s=256x256:d=0.1", "-c:v", "h264_nvenc", "-f", "null", "-"], capture_output=True)
        if probe.returncode == 0:
            return ["-c:v", "h264_nvenc", "-preset", "p7", "-rc", "vbr", "-cq", "21", "-b:v", "0", "-profile:v", "high"]
    return ["-c:v", "h264_mf", "-b:v", "12M"]


def photo_plan(spec):
    """Every photo on the drive as (relative path, time taken, source item, position in the clip), numbered like a camera roll."""
    shots = []
    for item in spec["photos"]:
        for i, frac in enumerate(item["stills"]):
            shots.append((dt.datetime.fromisoformat(item["date"]) + dt.timedelta(minutes=3 * i), item, frac))
    shots.sort(key=lambda s: s[0])
    return [(f"{item['folder']}/IMG_{n}.JPG", when, item, frac) for n, (when, item, frac) in enumerate(shots, start=6012)]


def build_photos(spec, dl: Downloads, out: Path, ff, probe):
    from PIL import Image
    for relative, when, item, frac in photo_plan(spec):
        src = dl.clip(item["clip"])
        target = out / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory() as tmp:
            frame = Path(tmp) / "f.png"
            run(ff, "-v", "error", "-y", "-ss", f"{duration(probe, src) * frac:.2f}", "-i", src, "-frames:v", "1", frame)
            im = Image.open(frame).convert("RGB")
        # Phone photos are 4:3 (3:4 when held upright); crop the clip's frame around its centre.
        w, h = im.size
        if w >= h:
            cw = min(w, round(h * 4 / 3))
            left = max(0, min(w - cw, round((w - cw) / 2 + item.get("shift", 0) * w)))
            im = im.crop((left, 0, left + cw, h))
        else:
            ch = min(h, round(w * 4 / 3))
            top = (h - ch) // 2
            im = im.crop((0, top, w, top + ch))
        exif = Image.Exif()
        exif[0x0132] = when.strftime("%Y:%m:%d %H:%M:%S")  # DateTime
        exif.get_ifd(0x8769)[0x9003] = when.strftime("%Y:%m:%d %H:%M:%S")  # DateTimeOriginal
        im.save(target, "JPEG", quality=92, exif=exif)
        stamp(target, when.isoformat())
        log(f"  {target.relative_to(out)}  <- mixkit {item['clip']} @ {frac:.2f}")


def build_videos(spec, dl: Downloads, out: Path, ff):
    enc = video_encoder(ff)
    for item in spec["videos"]:
        target = out / item["path"]
        target.parent.mkdir(parents=True, exist_ok=True)
        args, parts = [ff, "-v", "error", "-y"], []
        for k, (clip, start, end) in enumerate(item["segments"]):
            args += ["-ss", str(start), "-t", str(end - start), "-i", dl.clip(clip)]
            parts.append(f"[{k}:v]scale=1920:1080:force_original_aspect_ratio=increase,crop=1920:1080,fps=30,setsar=1,format=yuv420p[v{k}]")
        n = len(item["segments"])
        graph = ";".join(parts) + ";" + "".join(f"[v{k}]" for k in range(n)) + f"concat=n={n}:v=1:a=0[v]"
        # Phone videos carry a soundtrack; give each a quiet ambience so players behave as they would with real footage.
        total = sum(e - s for _, s, e in item["segments"])
        args += ["-f", "lavfi", "-t", str(total), "-i", "anoisesrc=color=brown:amplitude=0.02:sample_rate=48000",
                 "-filter_complex", graph + f";[{n}:a]lowpass=f=900,volume=0.6[a]", "-map", "[v]", "-map", "[a]",
                 *enc, "-c:a", "aac", "-b:a", "128k", "-movflags", "+faststart", target]
        run(*args)
        stamp(target, item["date"])
        log(f"  {target.relative_to(out)}  ({total} s)")


def build_memos(spec, dl: Downloads, out: Path, ff):
    import soundfile as sf
    from kokoro_onnx import Kokoro
    model = dl.get("kokoro/kokoro-v1.0.onnx", KOKORO + "kokoro-v1.0.onnx")
    voices = dl.get("kokoro/voices-v1.0.bin", KOKORO + "voices-v1.0.bin")
    tts = Kokoro(str(model), str(voices))
    for item in spec["memos"]:
        target = out / item["path"]
        target.parent.mkdir(parents=True, exist_ok=True)
        samples, rate = tts.create(item["text"], voice=item["voice"], speed=item.get("speed", 1.0), lang="en-us")
        with tempfile.TemporaryDirectory() as tmp:
            wav = Path(tmp) / "m.wav"
            sf.write(wav, samples, rate)
            # A phone memo: a breath of room tone either side, mono AAC at 48 kHz.
            run(ff, "-v", "error", "-y", "-i", wav, "-af", "adelay=600,apad=pad_dur=0.8,aresample=48000", "-ac", "1", "-c:a", "aac", "-b:a", "96k", target)
        stamp(target, item["date"])
        log(f"  {target.relative_to(out)}  ({item['voice']})")


def print_pdf(edge, html: Path, pdf: Path):
    with tempfile.TemporaryDirectory() as profile:
        run(edge, "--headless=new", "--disable-gpu", "--no-pdf-header-footer", f"--user-data-dir={profile}",
            f"--print-to-pdf={pdf}", html.as_uri())
    if not pdf.exists():
        sys.exit(f"{edge} did not print {html.name}. Set BROWSER to another Chromium browser.")


def scan(pdf: Path, target: Path, seed: int):
    """Turns a printed PDF into the image-only PDF a phone scanning app would make."""
    import pypdfium2 as pdfium
    from PIL import Image, ImageFilter, ImageEnhance, ImageChops
    rng = random.Random(seed)
    pages = []
    document = pdfium.PdfDocument(pdf.read_bytes())
    rendered = [document[i].render(scale=200 / 72).to_pil().convert("L") for i in range(len(document))]
    document.close()
    for im in rendered:
        im = im.rotate(rng.uniform(-0.6, 0.6), resample=Image.BICUBIC, expand=False, fillcolor=255)
        im = im.filter(ImageFilter.GaussianBlur(0.45))
        w, h = im.size
        # Uneven lighting and paper tone.
        shade = Image.radial_gradient("L").resize((w, h)).point(lambda v: 255 - v // 7)
        im = ImageChops.multiply(im, shade)
        noise = Image.effect_noise((w, h), 9).point(lambda v: 255 - abs(v - 128) // 3)
        im = ImageChops.multiply(im, noise)
        im = ImageEnhance.Contrast(im).enhance(1.08)
        tint = Image.merge("RGB", (im, im.point(lambda v: v * 0.99), im.point(lambda v: v * 0.955)))
        pages.append(tint)
    pages[0].save(target, "PDF", save_all=True, append_images=pages[1:], resolution=200, quality=82)


def build_documents(spec, out: Path, edge):
    from docx import Document
    from docx.shared import Pt
    from openpyxl import Workbook
    from openpyxl.styles import Font
    src = HERE / "documents"
    for n, item in enumerate(spec["pdf"]):
        target = out / item["path"]
        target.parent.mkdir(parents=True, exist_ok=True)
        if not edge:
            sys.exit("Chrome or Edge is needed to print the PDF documents.")
        if item.get("scan"):
            with tempfile.TemporaryDirectory() as tmp:
                printed = Path(tmp) / "printed.pdf"
                print_pdf(edge, src / item["html"], printed)
                scan(printed, target, seed=n)
        else:
            print_pdf(edge, src / item["html"], target)
        stamp(target, item["date"])
        log(f"  {target.relative_to(out)}{'  (scanned)' if item.get('scan') else ''}")
    for item in spec["docx"]:
        target = out / item["path"]
        target.parent.mkdir(parents=True, exist_ok=True)
        doc = Document()
        doc.styles["Normal"].font.name, doc.styles["Normal"].font.size = "Calibri", Pt(11)
        doc.add_heading(item["title"], level=1)
        for p in item["paragraphs"]:
            if p.startswith("## "):
                doc.add_heading(p[3:], level=2)
            elif p.startswith("- "):
                doc.add_paragraph(p[2:], style="List Bullet")
            else:
                doc.add_paragraph(p)
        doc.core_properties.author = "Sam Rivera"
        doc.save(target)
        stamp(target, item["date"])
        log(f"  {target.relative_to(out)}")
    for item in spec["xlsx"]:
        target = out / item["path"]
        target.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = item["sheet"]
        for row in item["rows"]:
            ws.append(row)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(12, max(len(str(c.value or "")) for c in col) + 2)
        wb.save(target)
        stamp(target, item["date"])
        log(f"  {target.relative_to(out)}")
    for item in spec["text"]:
        target = out / item["path"]
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(item["body"], encoding="utf-8")
        stamp(target, item["date"])
        log(f"  {target.relative_to(out)}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", required=True, type=Path, help="Demo drive root to create, e.g. ...\\FindItDemo\\Portable SSD")
    ap.add_argument("--cache", type=Path, default=Path(os.environ.get("USERPROFILE", ".")) / "FindItDemo" / "cache")
    ap.add_argument("--only", default=",".join(PARTS), help="Rebuild only these parts, in place: " + ", ".join(PARTS))
    ap.add_argument("--update-lock", action="store_true")
    ap.add_argument("--force", action="store_true", help="Replace an existing generated drive")
    a = ap.parse_args()
    out = a.out.resolve()
    parts = set(a.only.split(","))
    if parts - set(PARTS):
        sys.exit(f"Unknown part: {', '.join(parts - set(PARTS))}")
    marker = out / ".findit-demo-drive"
    if out.exists() and any(out.iterdir()):
        if not marker.exists():
            sys.exit(f"{out} exists and was not made by this script; refusing to touch it.")
        if not a.force:
            sys.exit(f"{out} already exists. Pass --force to rebuild it.")
        if (out / "FindIt").exists():
            sys.exit(f"{out} holds a prepared FindIt archive. Delete it by hand if you really want to rebuild.")
        for folder in {p for part in parts for p in PARTS[part]}:
            shutil.rmtree(out / folder, ignore_errors=True)
    out.mkdir(parents=True, exist_ok=True)
    marker.write_text("Generated by demo-video/library/build.py\n")
    ff, probe, edge = tools()
    dl = Downloads(a.cache, a.update_lock)
    try:
        if "photos" in parts or "videos" in parts:
            spec = json.loads((HERE / "sources.json").read_text(encoding="utf-8"))
            if "photos" in parts:
                log("Photos"); build_photos(spec, dl, out, ff, probe)
            if "videos" in parts:
                log("Videos"); build_videos(spec, dl, out, ff)
        if "memos" in parts:
            log("Voice memos"); build_memos(json.loads((HERE / "memos.json").read_text(encoding="utf-8")), dl, out, ff)
        if "documents" in parts:
            log("Documents"); build_documents(json.loads((HERE / "documents.json").read_text(encoding="utf-8")), out, edge)
    finally:
        dl.save()
    files = [p for p in out.rglob("*") if p.is_file() and p.name != ".findit-demo-drive"]
    log(f"Done: {len(files)} files, {sum(p.stat().st_size for p in files) / 1e6:.0f} MB in {out}")


if __name__ == "__main__":
    main()
